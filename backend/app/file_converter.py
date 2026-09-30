import csv
import re
from datetime import datetime, time, timedelta
from pathlib import Path

from openpyxl import load_workbook


TARGET_HEADERS = [
    "算法类别",
    "名称",
    "开始时机",
    "开始时间",
    "结束时机",
    "结束时间",
    "交易份数",
    "交易风格",
    "备注",
    "腿序号",
    "报单顺序",
    "合约代码",
    "交易编码",
    "委托比值",
    "买卖方向",
    "开平方向",
    "经纪商",
    "资产账号",
    "交易柜台",
    "柜台帐号",
    "切片个数",
]

REQUIRED_HEADERS = [
    "名称",
    "开始时间",
    "结束时间",
    "客户账号",
    "合约",
    "投机套保",
    "买卖方向",
    "开平方向",
    "数量（手）",
    "切片次数",
    "备注",
]

START_TIME_ADJUSTMENTS = {
    time(9, 30): time(9, 30, 5),
    time(13, 0): time(13, 0, 5),
}

END_TIME_ADJUSTMENTS = {
    time(11, 30): time(11, 29, 20),
    time(15, 0): time(14, 59, 20),
}

HEDGE_FLAG_MAP = {
    "投机": "Speculation",
    "套保": "Hedge",
}

SIDE_MAP = {
    "买": "Buy",
    "卖": "Sell",
}

OFFSET_MAP = {
    "开": "Open",
    "平": "Close",
}


def convert_xlsx_to_order_csv(xlsx_path, csv_path, exchange_resolver, now=None):
    source_path = Path(xlsx_path)
    target_path = Path(csv_path)
    now = now or datetime.now()

    try:
        _validate_source_file(source_path)
        workbook, worksheet = _load_worksheet(source_path)
        try:
            header_row = next(
                worksheet.iter_rows(min_row=1, max_row=1, values_only=True),
                None,
            )
            header_index = _build_header_index(header_row)
            _validate_required_headers(header_index)

            rows = []
            for instruction_index, row_data in _iter_order_rows(worksheet, header_index):
                rows.append(
                    _map_excel_row_to_csv_row(
                        row_data,
                        exchange_resolver,
                        now,
                        instruction_index,
                    )
                )

            _write_csv(target_path, TARGET_HEADERS, rows)
        finally:
            workbook.close()
        return True, str(target_path)
    except Exception as exc:
        return False, str(exc)


def _validate_source_file(xlsx_path):
    if not xlsx_path.exists():
        raise ValueError(f"文件不存在: {xlsx_path}")
    if xlsx_path.suffix.lower() != ".xlsx":
        raise ValueError(
            f"文件类型校验失败: 文件名={xlsx_path.name}, 期望文件类型=.xlsx"
        )


def _load_worksheet(xlsx_path):
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        return workbook, workbook.worksheets[0]
    except IndexError as exc:
        workbook.close()
        raise ValueError(f"Excel 中未找到工作表: {xlsx_path.name}") from exc


def _build_header_index(header_row):
    if not header_row:
        raise ValueError("Excel 表头为空")

    header_index = {}
    for index, header in enumerate(header_row):
        header_name = str(header).strip() if header is not None else ""
        if header_name:
            header_index[header_name] = index
    return header_index


def _validate_required_headers(header_index):
    missing_headers = [header for header in REQUIRED_HEADERS if header not in header_index]
    if missing_headers:
        raise ValueError(f"Excel 缺少必需表头: {', '.join(missing_headers)}")


def _iter_order_rows(worksheet, header_index):
    instruction_index = 0
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        if row is None:
            continue

        row_data = {}
        has_value = False
        for header, index in header_index.items():
            value = row[index] if index < len(row) else None
            row_data[header] = value
            if value not in (None, ""):
                has_value = True

        if has_value:
            instruction_index += 1
            yield instruction_index, row_data


def _map_excel_row_to_csv_row(row_data, exchange_resolver, now, instruction_index):
    start_time = row_data.get("开始时间")
    end_time = row_data.get("结束时间")

    start_at = "Immediately"
    start_value = ""
    schedule_date = None
    if not _is_empty_value(start_time):
        start_at = "CustomTime"
        start_value, schedule_date = _resolve_schedule_datetime(start_time, now, adjustments=START_TIME_ADJUSTMENTS)

    end_at = "Immediately"
    end_value = ""
    if not _is_empty_value(end_time):
        end_at = "CustomTime"
        end_value, _ = _resolve_schedule_datetime(end_time, now, base_date=schedule_date, adjustments=END_TIME_ADJUSTMENTS)

    return [
        "Twap",
        _stringify(row_data.get("名称")),
        start_at,
        start_value,
        end_at,
        end_value,
        _stringify(row_data.get("数量（手）")),
        "0000000002",
        _stringify(row_data.get("备注")),
        "",
        "",
        _format_contract_code(
            row_data.get("合约"),
            exchange_resolver,
            instruction_index,
        ),
        _map_enum_value(row_data.get("投机套保"), HEDGE_FLAG_MAP, "投机套保"),
        "",
        _map_enum_value(row_data.get("买卖方向"), SIDE_MAP, "买卖方向"),
        _map_enum_value(row_data.get("开平方向"), OFFSET_MAP, "开平方向"),
        "gtjaqh",
        _stringify(row_data.get("客户账号")),
        "ctp_prod",
        "1002156",
        _stringify(row_data.get("切片次数")),
    ]


def _resolve_schedule_datetime(cell_value, now, base_date=None, adjustments=None):
    cell_time = _extract_time_value(cell_value)
    if adjustments and cell_time in adjustments:
        cell_time = adjustments[cell_time]
    target_date = base_date or _resolve_schedule_date(cell_time, now)
    target_datetime = datetime.combine(target_date, cell_time)
    return target_datetime.strftime("%Y-%m-%d %H:%M:%S.000+0800"), target_date


def _resolve_schedule_date(cell_time, now):
    current_date = now.date()
    if now.time() > cell_time:
        return current_date + timedelta(days=1)
    return current_date


def _format_contract_code(contract, exchange_resolver, instruction_index):
    contract_value = _stringify(contract)
    if not contract_value:
        raise ValueError(f"第{instruction_index}个指令合约信息错误: 合约不能为空")

    exchange = exchange_resolver(contract_value)
    exchange_value = _stringify(exchange).lower()
    if not exchange_value:
        raise ValueError(
            f"第{instruction_index}个指令合约信息错误: 未找到合约对应交易所, 合约={contract_value}"
        )
    return f"{contract_value.lower()}.{exchange_value}"


def _write_csv(csv_path, header, rows):
    with csv_path.open("w", encoding="utf-8", newline="") as csv_file:
        writer = csv.writer(csv_file)
        writer.writerow(header)
        writer.writerows(rows)


def _extract_time_value(cell_value):
    if isinstance(cell_value, datetime):
        return cell_value.time().replace(microsecond=0)
    if isinstance(cell_value, time):
        return cell_value.replace(microsecond=0)
    if isinstance(cell_value, str):
        text = cell_value.strip()
        if not text:
            raise ValueError("时间字段为空字符串")
        text = re.sub(r"\s+", " ", text)
        for fmt in ("%Y/%m/%d %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%H:%M:%S", "%m %d %Y %I:%M:%S %p"):
            try:
                return datetime.strptime(text, fmt).time()
            except ValueError:
                continue
    raise ValueError(f"无法识别时间字段格式: {cell_value}")


def _map_enum_value(value, mapping, field_name):
    key = _stringify(value)
    if key not in mapping:
        raise ValueError(f"{field_name}字段值不支持: {key}")
    return mapping[key]


def _stringify(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _is_empty_value(value):
    if value is None:
        return True
    if isinstance(value, str):
        return not value.strip()
    return False
